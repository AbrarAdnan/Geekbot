"""CSV, TSV and Excel loader: a schema card per table, never the rows (step S1.1, D-051).

**Why a card and not the rows.** Chat mode answers by retrieving text and asking a model to read
it. Asked "what is the total revenue?", it would add up whichever rows happened to be retrieved
and state the sum with confidence, which is exactly the guessing CSV mode exists to avoid (there
the model writes one SQL query, the query runs on the real data, and the numbers in the answer
are checked against the result). So chat mode indexes only what a table *is*: its columns, their
types and a few example values. That is enough to find the right file for a question, and the
card itself tells the reader to use CSV mode for any figure. (This replaces D-016's "row text for
files up to 2,000 rows".)

No pandas here: the loader streams the file with the standard ``csv`` module (and ``openpyxl`` in
read-only mode for Excel), so a big file costs time but not memory, and importing Geekbot stays
fast. Types, empty counts and examples come from the first ``SAMPLE_ROWS`` rows; only the row
count looks at the whole file.
"""

import csv
import datetime
import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Final

from geekbot.ingest.loaders.common import detect_encoding
from geekbot.ingest.loaders.model import LoadedDocument, LoadError, Segment

TABULAR_EXTENSIONS: Final = frozenset({".csv", ".tsv", ".xlsx"})

SAMPLE_ROWS: Final = 1000  # rows used for types, empty counts and examples
MAX_COLUMNS_SHOWN: Final = 60
EXAMPLES_PER_COLUMN: Final = 3
MAX_EXAMPLE_CHARS: Final = 30

_ENCODING_SNIFF_BYTES: Final = 100_000
_DELIMITER_SNIFF_CHARS: Final = 16_384
_INTEGER: Final = re.compile(r"[+-]?(0|[1-9]\d*)")
_DECIMAL: Final = re.compile(r"[+-]?(\d+\.\d*|\.\d+)([eE][+-]?\d+)?")
# 2024-03, 2024-03-15, 2024-03-15 08:30[:00], 2024-03-15T08:30[:00][Z]. Nothing looser: a
# date format that guesses is a type that lies.
_DATE: Final = re.compile(r"\d{4}-\d{2}(-\d{2}([ T]\d{2}:\d{2}(:\d{2})?Z?)?)?")
# Tried in this order: "1" is an integer before it could be anything else.
_TEXT_PATTERNS: Final = (("integer", _INTEGER), ("decimal", _DECIMAL), ("date", _DATE))


@dataclass(frozen=True)
class _Table:
    """What the card needs from one sheet or file."""

    name: str | None  # the sheet name for Excel, None for CSV and TSV
    header: list[str]
    sample: list[list[object]]  # up to SAMPLE_ROWS data rows, each as wide as the header
    row_count: int  # every non-blank data row in the file


def load_table(path: Path) -> LoadedDocument:
    """One schema-card segment per table: one for a CSV or TSV, one per visible Excel sheet."""
    extension = path.suffix.lower()
    if extension == ".xlsx":
        tables = _read_xlsx(path)
        label = "Excel sheet"
    else:
        tables = [_read_delimited(path, "\t" if extension == ".tsv" else None)]
        label = "TSV file" if extension == ".tsv" else "CSV file"
    segments = tuple(Segment(_card(path.name, label, table)) for table in tables)
    return LoadedDocument(path=path, kind="table", title=path.name, segments=segments)


# --- reading ---------------------------------------------------------------------------------


def _read_delimited(path: Path, delimiter: str | None) -> _Table:
    """Stream a CSV or TSV. ``delimiter=None`` sniffs it from the first lines (default ``,``)."""
    with path.open("rb") as raw:
        head = raw.read(_ENCODING_SNIFF_BYTES)
    encoding = detect_encoding(head)
    if delimiter is None:
        sniffed = head.decode(encoding, errors="replace")[:_DELIMITER_SNIFF_CHARS]
        delimiter = _sniff_delimiter(sniffed)
    # errors="replace": a stray bad byte far into a big file turns into U+FFFD in one example
    # value instead of failing the whole file; this is a description of the table, not its data.
    with path.open(newline="", encoding=encoding, errors="replace") as handle:
        table = _collect(csv.reader(handle, delimiter=delimiter), name=None)
    if table is None:
        raise LoadError("the file is empty")
    return table


def _sniff_delimiter(text: str) -> str:
    try:
        return csv.Sniffer().sniff(text, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _read_xlsx(path: Path) -> list[_Table]:
    """Stream every visible sheet of a workbook, reading cached formula results (not formulas)."""
    from openpyxl import load_workbook  # type: ignore[import-untyped, unused-ignore]

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        tables: list[_Table] = []
        for sheet in workbook.worksheets:
            if getattr(sheet, "sheet_state", "visible") != "visible":
                continue  # hidden on purpose, so not described either
            table = _collect(sheet.iter_rows(values_only=True), name=str(sheet.title))
            if table is not None:
                tables.append(table)
    finally:
        workbook.close()
    if not tables:
        raise LoadError("the workbook has no visible sheet with data")
    return tables


def _collect(rows: Iterable[Sequence[object]], *, name: str | None) -> _Table | None:
    """Read the header (the first non-blank row) and sample the rest. ``None`` if no header."""
    iterator = iter(rows)
    header_row = next((row for row in iterator if not _is_blank(row)), None)
    if header_row is None:
        return None
    header = [_cell_text(cell) for cell in header_row]
    while header and not header[-1]:  # Excel pads rows to the sheet width
        header.pop()
    if not header:
        return None
    header = [text or f"column {number}" for number, text in enumerate(header, start=1)]
    width = len(header)
    sample: list[list[object]] = []
    count = 0
    for row in iterator:
        if _is_blank(row):
            continue
        count += 1
        if len(sample) < SAMPLE_ROWS:
            sample.append([row[index] if index < len(row) else None for index in range(width)])
    return _Table(name=name, header=header, sample=sample, row_count=count)


def _is_blank(row: Sequence[object]) -> bool:
    return all(_is_empty(cell) for cell in row)


def _is_empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _cell_text(value: object) -> str:
    return "" if value is None else " ".join(str(value).split())


# --- describing ------------------------------------------------------------------------------


def _classify(value: object) -> str:
    """The type of one non-empty value: boolean, integer, decimal, date or text.

    Excel hands over real Python values (``int``, ``float``, ``datetime``); a CSV hands over
    text, so text is matched against the patterns above.
    """
    if isinstance(value, bool):  # before int: True is an int in Python
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "decimal"
    if isinstance(value, datetime.date | datetime.time):
        return "date"
    return _classify_text(str(value).strip())


def _classify_text(text: str) -> str:
    if text.lower() in {"true", "false"}:
        return "boolean"
    for kind, pattern in _TEXT_PATTERNS:
        if pattern.fullmatch(text):
            return kind
    return "text"


def _column_type(values: list[object]) -> str:
    """One type for a column: its values' common type; integers mixed with decimals are
    decimal; any other mix is text; a column with no values is ``empty``."""
    kinds = {_classify(value) for value in values if not _is_empty(value)}
    if not kinds:
        return "empty"
    if len(kinds) == 1:
        return next(iter(kinds))
    if kinds <= {"integer", "decimal"}:
        return "decimal"
    return "text"


def _examples(values: list[object]) -> list[str]:
    """The first few distinct non-empty values, shortened and safe to put in a Markdown cell."""
    found: list[str] = []
    for value in values:
        if _is_empty(value):
            continue
        text = _cell_text(value)
        if len(text) > MAX_EXAMPLE_CHARS:
            text = text[: MAX_EXAMPLE_CHARS - 3] + "..."
        if text not in found:
            found.append(text)
            if len(found) == EXAMPLES_PER_COLUMN:
                break
    return found


def _md(text: str) -> str:
    """Make text safe inside a Markdown table cell (a ``|`` would split the column)."""
    return text.replace("\\", "\\\\").replace("|", "\\|")


def _card(filename: str, label: str, table: _Table) -> str:
    """The Markdown description of one table. Deterministic: same file, same text, same hash."""
    where = f"{filename}, sheet {table.name}" if table.name is not None else filename
    columns = len(table.header)
    scope = (
        f"all {table.row_count:,} rows"
        if len(table.sample) == table.row_count
        else f"the first {len(table.sample):,} rows"
    )
    lines = [
        f"# Table: {_md(where)}",
        "",
        (
            f"{label} with {table.row_count:,} rows and {columns} columns. "
            f"Column types, empty shares and example values come from {scope}."
        ),
        "",
        "| column | type | empty | examples |",
        "|---|---|---|---|",
    ]
    for index, column in enumerate(table.header[:MAX_COLUMNS_SHOWN]):
        values = [row[index] for row in table.sample]
        empty = sum(_is_empty(value) for value in values)
        share = f"{round(100 * empty / len(values))}%" if values else "-"
        examples = ", ".join(_examples(values))
        lines.append(f"| {_md(column)} | {_column_type(values)} | {share} | {_md(examples)} |")
    if columns > MAX_COLUMNS_SHOWN:
        lines.append("")
        lines.append(f"{columns - MAX_COLUMNS_SHOWN} more columns are not described here.")
    lines.append("")
    lines.append(
        "This card describes the shape of the table only: chat mode does not index its rows. "
        "For counts, totals, averages or any other figure, open the file in CSV mode."
    )
    return "\n".join(lines)
