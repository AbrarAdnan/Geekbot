"""Write the Word, PowerPoint and Excel files of the golden corpus (golden v0, D-067).

Run once, from the repository root, in the ``geekbot`` conda env::

    python evals/corpus_src/build_office_files.py

It writes four files into ``evals/corpus/`` and one Word file next to this script::

    evals/corpus/policy/ai-use-policy.docx
    evals/corpus/policy/onboarding-letter.docx
    evals/corpus/project/rollout-kickoff.pptx
    evals/corpus/project/tools-and-costs.xlsx
    evals/corpus_src/vendor-memo.docx        (not part of the corpus: see below)

**The PDF is made by hand.** The corpus also holds ``project/vendor-memo.pdf``. No PDF library is
installed, so open ``vendor-memo.docx`` in Word and use "Save as PDF" to write
``evals/corpus/project/vendor-memo.pdf``. The Word file stays here, outside the folder that
``geekbot ingest`` reads, because it would otherwise be a second copy of the same memo.

**Why a script and not the files alone.** The facts of the fictional Brightwater Studio (names,
dates, prices) recur across all fifteen documents. Keeping the text of the binary ones here, in
one place, makes them reviewable as text and re-creatable; the generated files are committed with
the corpus. The test ``tests/unit/test_golden_corpus.py`` builds them again, reads them back
through Geekbot's own loaders and checks every golden question against what they say.

Everything is original work and carries the repository's license.
"""

import argparse
import sys
from collections.abc import Callable, Sequence
from pathlib import Path
from typing import Any

SOURCE_DIR = Path(__file__).resolve().parent
CORPUS_DIR = SOURCE_DIR.parent / "corpus"

MEMO_DOCX_NAME = "vendor-memo.docx"
MEMO_PDF_PATH = "project/vendor-memo.pdf"

Table = Sequence[Sequence[str]]  # the first row is the header

# --- Word: the AI use policy -----------------------------------------------------------------

POLICY_TITLE = "Brightwater Studio: AI use policy"
POLICY_INTRO = (
    "Version 1.0. Approved by Maya Okafor, managing director, on 21 March 2025. "
    "Owner: Daniel Reyes. Next review due: 21 September 2025."
)
POLICY_DATA_CLASSES: Table = (
    ("Class", "What it covers", "May go into Lantern?", "May go into outside AI tools?"),
    (
        "Green",
        "Published or public material, such as marketing copy already on the website",
        "Yes",
        "Yes",
    ),
    ("Amber", "Client work in progress and internal plans", "Yes", "No"),
    (
        "Red",
        "Client contracts, personal data and financial records",
        "Only with written approval from Tomasz Lewandowski",
        "No",
    ),
)
POLICY_RULES = (
    "Use only your own account. Sharing a sign-in is not allowed.",
    "Check the workspace name at the top of a chat before you paste client material.",
    (
        "A person must read and check every answer before it leaves the studio. Lantern drafts, "
        "and the person who sends the answer is responsible for it."
    ),
    "Do not use outside cloud AI tools for amber or red data.",
    "Attend the 45 minute training session before an account is issued.",
)
POLICY_RETENTION: Table = (
    ("What", "How long it is kept"),
    ("Chat logs", "90 days"),
    ("Audit logs", "365 days (one year)"),
    ("Uploaded documents", "until the project is closed, plus 12 months"),
    ("Backups", "30 days"),
)

# --- Word: the welcome letter ----------------------------------------------------------------

LETTER_TITLE = "Welcome to Lantern"
LETTER_PARAGRAPHS = (
    "From Daniel Reyes, operations lead, to everyone joining Lantern.",
    (
        "Lantern is the private AI assistant of Brightwater Studio. It runs on a computer in our "
        "own comms cupboard, so what you give it stays inside the building. You reach it from "
        "the office network at lantern.office.lan; it does not work from home."
    ),
    (
        "Before you can sign in you must attend one training session. It takes 45 minutes. The "
        "first sessions were held on 5 and 6 May 2025; people who join later book a session "
        "with me. Once you have attended, I create your account, usually within one working day."
    ),
)
LETTER_FIRST_WEEK = (
    (
        "Try it on something small first, such as a summary of last week's meeting notes. "
        "Drafting an email to a client and searching a project's documents are the other two "
        "things most people use it for."
    ),
    (
        "Look at the workspace name at the top of every chat. One workspace belongs to one "
        "client project, and the name is there so that nothing lands in the wrong place."
    ),
    (
        "Remember the three kinds of data. Green is public material, amber is client work in "
        "progress, and red is contracts, personal data and money. Red data does not go in "
        "without written approval from Tomasz Lewandowski."
    ),
)
LETTER_CONTACTS = (
    "Questions about using Lantern: Daniel Reyes.",
    "Lantern is slow or broken: Priya Nair, who is on call on Mondays and Thursdays.",
    "Something may have gone wrong with data: Tomasz Lewandowski, within 24 hours.",
)
LETTER_CLOSING = "After your first four weeks I will ask how it is going. Welcome aboard."

# --- Word: the vendor memo (source of the PDF) -----------------------------------------------

MEMO_TITLE = "Memo: Which way should Brightwater run an AI assistant?"
MEMO_HEADER = (
    "To: Maya Okafor",
    "From: Daniel Reyes",
    "Date: 28 February 2025",
    "Decision wanted: at the kickoff meeting on 3 March 2025",
)
MEMO_RECOMMENDATION = (
    "I recommend Option C, a computer of our own running the assistant in the office. It is the "
    "most expensive of the three, but it is the only one that keeps client work inside the "
    "building, and two of our clients require that."
)
MEMO_OPTIONS = (
    (
        "Option A: a cloud subscription",
        (
            "Each person pays £22 per user per month. For 12 people that is £264 per month, or "
            "£3,168 for the year. Nothing has to be bought or looked after. Our prompts and "
            "documents would be processed by the provider, in an EU region."
        ),
    ),
    (
        "Option B: a hosted private instance",
        (
            "A provider in the UK runs an assistant for us alone at a flat £780 per month, or "
            "£9,360 for the year. Our data stays with that provider. Corvid Foods would have to "
            "agree to a change in its NDA first."
        ),
    ),
    (
        "Option C: a workstation in our own office",
        (
            "A one-off cost of £2,330 (£2,150 for the workstation and £180 for a UPS battery "
            "unit) and £1,244 per month (a support retainer of £1,200, about £35 of electricity "
            "and £9 of cloud storage for backups). The first year costs £17,258. The second year "
            "would cost £14,928, because the one-off cost does not come back. Data never leaves "
            "the building."
        ),
    ),
)
MEMO_COMPARISON: Table = (
    ("Option", "One-off", "Per month", "First year", "Where the data goes"),
    ("A: cloud subscription", "£0", "£264", "£3,168", "Provider, EU region"),
    ("B: hosted private instance", "£0", "£780", "£9,360", "Provider, UK"),
    ("C: our own workstation", "£2,330", "£1,244", "£17,258", "Stays in our office"),
)
MEMO_WHAT_MATTERS = (
    (
        "Two clients, Harbour and Pine and Corvid Foods, have NDAs that do not allow third "
        "parties to process their work. Option A would break both. Option B would break the "
        "Corvid Foods NDA unless it were changed."
    ),
    (
        "Option C costs £14,090 more than Option A in the first year. We accept that as the "
        "price of being able to take the work of both clients."
    ),
    (
        "Speed target: the 95th percentile of the time to the first word of an answer should be "
        "under 3 seconds. We will measure it before the pilot."
    ),
    (
        "Risk of Option C: one machine is a single point of failure. A nightly backup and a UPS "
        "battery unit reduce the risk, and Priya Nair will write down how to rebuild it."
    ),
)

# --- PowerPoint: the kickoff deck ------------------------------------------------------------

DECK_TITLE = "Project Lantern: kickoff"
DECK_SUBTITLE = "Brightwater Studio, 3 March 2025. Presenter: Daniel Reyes"
DECK_SLIDES: Sequence[tuple[str, Sequence[str]]] = (
    (
        "Why we are doing this",
        (
            (
                "Two clients, Harbour and Pine and Corvid Foods, have NDAs that forbid third "
                "parties from processing their work"
            ),
            "Staff already try outside AI tools on their own",
            "Goal: a private assistant that keeps client work inside the building",
        ),
    ),
    (
        "What Lantern will do",
        (
            "Summarise meeting notes",
            "Draft client emails",
            "Search project documents",
            "It will not decide anything about people",
        ),
    ),
    (
        "Options we compared",
        (
            "A: cloud subscription, £22 per user per month, £3,168 a year for 12 people",
            "B: hosted private instance, £780 per month, £9,360 a year",
            "C: our own workstation, £17,258 in the first year",
            "Recommended: C, because it is the only option that suits both NDAs",
        ),
    ),
    (
        "Who does what",
        (
            "Maya Okafor: sponsor",
            "Daniel Reyes: project owner",
            "Priya Nair: builds and runs the workstation",
            "Tomasz Lewandowski: policy and security",
            "Hannah Frost: leads the pilot",
            "Samir Haddad: budget",
        ),
    ),
    (
        "Timeline",
        (
            "17 March: workstation ordered",
            "21 March: AI use policy approved",
            "27 March: hardware delivered; 28 to 31 March: set-up and first benchmark",
            "7 April to 2 May: four-week pilot with four people",
            "5 and 6 May: training for all staff",
            "12 May: rollout to all 12 staff",
            "21 September: first policy review",
        ),
    ),
    (
        "Budget",
        (
            "One-off: £2,330 (workstation £2,150, UPS £180)",
            "Every month: £1,244 (support retainer £1,200, electricity £35, backup storage £9)",
            "First year in total: £17,258",
        ),
    ),
    (
        "Risks and what we do about them",
        (
            "Someone pastes red data: policy and training before anyone gets an account",
            "One machine can fail: nightly backup and a UPS battery unit",
            "Too many people at once: three can chat together, a fourth waits",
            "Knowledge sits with one person: Priya Nair writes down how to rebuild it",
        ),
    ),
)

# --- Excel: tools and costs ------------------------------------------------------------------

TOOLS_SHEET = "Tools"
TOOLS_ROWS: Table = (
    ("tool", "purpose", "license", "owner", "monthly_cost_gbp", "status"),
    ("Inference server", "runs the chat model", "open source", "Priya Nair", "0", "live"),
    ("Web chat interface", "the screen staff use", "open source", "Priya Nair", "0", "live"),
    (
        "Vector store",
        "search index for uploaded documents",
        "open source",
        "Priya Nair",
        "0",
        "live",
    ),
    ("Backup agent", "nightly backup to the NAS", "open source", "Priya Nair", "0", "live"),
    ("Retention job", "deletes expired logs", "in-house", "Daniel Reyes", "0", "live"),
    ("Cloud storage", "weekly encrypted backup copy", "commercial", "Priya Nair", "9", "live"),
    ("Support retainer", "up to 4 days a month", "contract", "Daniel Reyes", "1200", "live"),
    ("Electricity", "running costs of lantern-01", "utility", "Samir Haddad", "35", "live"),
)
BUDGET_SHEET = "Budget"
BUDGET_ROWS: Table = (
    ("item", "type", "amount_gbp", "date_paid"),
    ("Workstation lantern-01", "one-off", "2150", "2025-03-17"),
    ("UPS battery unit", "one-off", "180", "2025-03-17"),
    ("Support retainer", "monthly", "1200", ""),
    ("Electricity", "monthly", "35", ""),
    ("Cloud backup storage", "monthly", "9", ""),
)


# --- builders --------------------------------------------------------------------------------


def _numeric(text: str) -> int | str:
    """Excel should hold amounts as numbers, so a column of digits becomes integers."""
    return int(text) if text.isdigit() else text


def _add_table(document: Any, rows: Table) -> None:
    """A bordered table whose first row is the header (the document is a ``docx.Document``)."""
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, text in enumerate(row):
            table.cell(row_index, column_index).text = text
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True


def build_policy(path: Path) -> None:
    """The AI use policy as a Word file."""
    from docx import Document

    document = Document()
    document.add_heading(POLICY_TITLE, level=0)
    document.add_paragraph(POLICY_INTRO)
    document.add_heading("1. Purpose and scope", level=1)
    document.add_paragraph(
        "This policy sets the rules for using artificial intelligence tools at Brightwater "
        "Studio. It applies to all 12 members of staff and to contractors. It covers Lantern, "
        "the studio's private assistant, and any other AI tool, whether or not the studio pays "
        "for it."
    )
    document.add_heading("2. Data classes", level=1)
    document.add_paragraph("Every piece of information belongs to one of three classes.")
    _add_table(document, POLICY_DATA_CLASSES)
    document.add_heading("3. Rules for everyone", level=1)
    for rule in POLICY_RULES:
        document.add_paragraph(rule, style="List Bullet")
    document.add_heading("4. Retention", level=1)
    _add_table(document, POLICY_RETENTION)
    document.add_paragraph(
        "Retention is enforced automatically each night, except where a legal hold is in place. "
        "A legal hold is placed by the person handling an incident and is removed only by Tomasz "
        "Lewandowski."
    )
    document.add_heading("5. Incidents", level=1)
    document.add_paragraph(
        "Report any suspected incident to Tomasz Lewandowski within 24 hours. If red data may "
        "have been exposed, phone him within 1 hour. The incident runbook says what happens next."
    )
    document.add_heading("6. Approval and review", level=1)
    document.add_paragraph(
        "Maya Okafor approved this policy on 21 March 2025. It is reviewed every six months, so "
        "the first review is due on 21 September 2025."
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def build_onboarding_letter(path: Path) -> None:
    """The welcome letter for new users as a Word file."""
    from docx import Document

    document = Document()
    document.add_heading(LETTER_TITLE, level=0)
    for paragraph in LETTER_PARAGRAPHS:
        document.add_paragraph(paragraph)
    document.add_heading("Your first week", level=1)
    for paragraph in LETTER_FIRST_WEEK:
        document.add_paragraph(paragraph, style="List Bullet")
    document.add_heading("Who to ask", level=1)
    for paragraph in LETTER_CONTACTS:
        document.add_paragraph(paragraph, style="List Bullet")
    document.add_paragraph(LETTER_CLOSING)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def build_vendor_memo(path: Path) -> None:
    """The vendor memo as a Word file; the PDF in the corpus is its "Save as PDF"."""
    from docx import Document

    document = Document()
    document.add_heading(MEMO_TITLE, level=0)
    for line in MEMO_HEADER:
        document.add_paragraph(line)
    document.add_heading("Recommendation", level=1)
    document.add_paragraph(MEMO_RECOMMENDATION)
    document.add_heading("The three options", level=1)
    for heading, text in MEMO_OPTIONS:
        document.add_heading(heading, level=2)
        document.add_paragraph(text)
    document.add_heading("Comparison", level=1)
    _add_table(document, MEMO_COMPARISON)
    document.add_heading("What matters most", level=1)
    for text in MEMO_WHAT_MATTERS:
        document.add_paragraph(text, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def build_kickoff_deck(path: Path) -> None:
    """The kickoff slides as a PowerPoint file: a title slide, then one slide per topic."""
    from pptx import Presentation

    deck = Presentation()
    title_slide = deck.slides.add_slide(deck.slide_layouts[0])
    title_slide.shapes.title.text = DECK_TITLE
    title_slide.placeholders[1].text = DECK_SUBTITLE
    for title, bullets in DECK_SLIDES:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = bullets[0]
        for bullet in bullets[1:]:
            frame.add_paragraph().text = bullet
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(path))


def build_tools_workbook(path: Path) -> None:
    """Two sheets, Tools and Budget. Chat mode indexes only their shape (D-051)."""
    from openpyxl import Workbook

    book = Workbook()
    sheets = ((TOOLS_SHEET, TOOLS_ROWS), (BUDGET_SHEET, BUDGET_ROWS))
    for index, (name, rows) in enumerate(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = name
        sheet.append(list(rows[0]))
        for row in rows[1:]:
            sheet.append([_numeric(cell) if cell else None for cell in row])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(path))


# The files that go into the corpus, by their path inside it. The test reads this manifest.
OFFICE_FILES: dict[str, Callable[[Path], None]] = {
    "policy/ai-use-policy.docx": build_policy,
    "policy/onboarding-letter.docx": build_onboarding_letter,
    "project/rollout-kickoff.pptx": build_kickoff_deck,
    "project/tools-and-costs.xlsx": build_tools_workbook,
}


def build_all(corpus_dir: Path = CORPUS_DIR, source_dir: Path = SOURCE_DIR) -> list[Path]:
    """Write every office file into ``corpus_dir`` and the memo's Word file into ``source_dir``."""
    written: list[Path] = []
    for relative, builder in OFFICE_FILES.items():
        target = corpus_dir / relative
        builder(target)
        written.append(target)
    memo = source_dir / MEMO_DOCX_NAME
    build_vendor_memo(memo)
    written.append(memo)
    return written


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0] if __doc__ else None)
    parser.add_argument("--corpus", type=Path, default=CORPUS_DIR, help="the corpus folder")
    parser.add_argument("--source", type=Path, default=SOURCE_DIR, help="where the memo goes")
    args = parser.parse_args(argv)
    for path in build_all(args.corpus, args.source):
        sys.stdout.write(f"wrote {path}\n")
    sys.stdout.write(
        f"\nNow open {args.source / MEMO_DOCX_NAME} in Word and save it as PDF to "
        f"{args.corpus / MEMO_PDF_PATH}\n"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
